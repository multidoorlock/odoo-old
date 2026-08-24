#!/usr/bin/env python3
"""Reproducible integrity checks for the generated MasterProducts catalog."""

import argparse
import base64
import gzip
import json
import math
import re
import subprocess
import sys
import tempfile
from pathlib import Path


FILTER_GROUP = "40 - מערכת סינון"
FILTER_TEMPLATE_COUNTS = {
    "40 - מערכת סינון — רב בריח — לביא PRO": 1,
    "40 - מערכת סינון — רב בריח — כפיר+": 11,
    "40 - מערכת סינון — בית אל — Rainbow": 1,
    "40 - מערכת סינון — בית אל — Hidden": 5,
}
WHITE_WING_TEMPLATE_COUNTS = {
    '1201 - כנף לדלת ממ"ד לבן — פתיחה רגילה': 22,
    '1201 - כנף לדלת ממ"ד לבן — הזזה': 16,
}
TOKEN_RE = re.compile(r"\[([^\[\]]+)\]")
DIRECTION_MARKER_RE = re.compile(
    r"(?<![A-Za-z])([LRD])(?![A-Za-z])", re.IGNORECASE
)
HEBREW_RE = re.compile(r"[\u0590-\u05ff]")
BASE_NAME_TOKENS = {"שם קבוצת פריטים", "קבוצת פריטים", "דגם", "שם דגם"}


def clean(value):
    return " ".join(str(value or "").split())


def normalize_token(value):
    return clean(value).casefold()


def split_direction_marker(value):
    value = clean(value)
    markers = DIRECTION_MARKER_RE.findall(value)
    if not HEBREW_RE.search(value) or len(markers) != 1:
        return value, ""
    return clean(DIRECTION_MARKER_RE.sub(" ", value)), markers[0].upper()


def split_name_format(format_value, attribute_names):
    attributes_by_key = {
        normalize_token(name): clean(name) for name in attribute_names
    }
    base_keys = {normalize_token(name) for name in BASE_NAME_TOKENS}
    recognized = []
    for match in TOKEN_RE.finditer(format_value or ""):
        key = normalize_token(match.group(1))
        if key in base_keys or key in attributes_by_key:
            recognized.append((match, key))
    base_end = max(
        (match.end() for match, key in recognized if key in base_keys),
        default=0,
    )
    attribute_matches = [
        (match, key)
        for match, key in recognized
        if key in attributes_by_key
    ]
    rules = {
        normalize_token(name): {
            "mode": "hidden",
            "prefix": "",
            "suffix": "",
        }
        for name in attribute_names
    }
    previous_match = None
    previous_key = None
    for match, key in attribute_matches:
        if previous_match is None:
            rules[key]["prefix"] = (format_value or "")[base_end:match.start()]
        else:
            rules[previous_key]["suffix"] = (format_value or "")[
                previous_match.end():match.start()
            ]
        rules[key]["mode"] = "value"
        previous_match = match
        previous_key = key
    final_suffix = (
        (format_value or "")[previous_match.end():]
        if previous_match is not None
        else (format_value or "")[base_end:]
    )
    return rules, final_suffix


def name_with_group(group_name, model_name):
    group_name = clean(group_name)
    model_name = clean(model_name)
    if model_name == group_name:
        model_name = ""
    elif group_name and model_name.startswith(f"{group_name} "):
        model_name = model_name[len(group_name) + 1:]
    return clean(f"{group_name} {model_name}")


def fail(message):
    raise AssertionError(message)


def load_catalog(json_path):
    data = json.loads(json_path.read_text(encoding="utf-8"))
    payload_path = json_path.with_name("catalog.json.gz.b64")
    decoded = gzip.decompress(
        base64.b64decode(payload_path.read_text(encoding="ascii"))
    ).decode("utf-8")
    if json.loads(decoded) != data:
        fail("The compressed payload does not match catalog.json")
    return data


def validate_references(data):
    groups = {item["key"]: item for item in data["groups"]}
    attributes = {item["key"]: item for item in data["attributes"]}
    values = {item["key"]: item for item in data["values"]}
    if len(values) != len(data["values"]):
        fail("Duplicate global value keys")

    all_skus = []
    for template in data["templates"]:
        if template["group_key"] not in groups:
            fail(f"Unknown group in {template['key']}")
        lines = template["attribute_lines"]
        line_attributes = [line["attribute_key"] for line in lines]
        if len(line_attributes) != len(set(line_attributes)):
            fail(f"Duplicate attribute line in {template['key']}")
        if [line["sequence"] for line in lines] != sorted(
            line["sequence"] for line in lines
        ):
            fail(f"Attribute lines are out of order in {template['key']}")

        allowed_by_attribute = {}
        for line in lines:
            attribute_key = line["attribute_key"]
            if attribute_key not in attributes:
                fail(f"Unknown attribute {attribute_key} in {template['key']}")
            allowed = []
            for value_data in line["values"]:
                value_key = value_data["value_key"]
                if value_key not in values:
                    fail(f"Unknown value {value_key} in {template['key']}")
                if values[value_key]["attribute_key"] != attribute_key:
                    fail(f"Value {value_key} belongs to the wrong attribute")
                allowed.append(value_key)
            if len(allowed) != len(set(allowed)):
                fail(f"Duplicate line value in {template['key']}")
            allowed_by_attribute[attribute_key] = set(allowed)

        combinations = set()
        for variant in template["variants"]:
            all_skus.append(variant["sku"])
            if len(variant["value_keys"]) != len(lines):
                fail(f"Wrong value count for {variant['sku']}")
            for line, value_key in zip(lines, variant["value_keys"]):
                if value_key not in allowed_by_attribute[line["attribute_key"]]:
                    fail(f"Value {value_key} is not assigned to {template['key']}")
            combination = tuple(variant["value_keys"])
            if combination in combinations:
                fail(f"Duplicate combination in {template['key']}")
            combinations.add(combination)

        expected_cartesian = math.prod(len(line["values"]) for line in lines)
        if template["cartesian_count"] != expected_cartesian:
            fail(f"Wrong Cartesian count in {template['key']}")
        if template["allowed_count_after_exclusions"] < len(combinations):
            fail(f"Pair exclusions remove a valid combination in {template['key']}")

    all_skus.extend(item["sku"] for item in data["unique_items"])
    if len(all_skus) != data["metadata"]["total_products"]:
        fail("Product total does not match metadata")
    if len(all_skus) != len(set(all_skus)):
        fail("Duplicate final SKUs")
    return set(all_skus)


def validate_rendered_output(data):
    groups = {item["key"]: item for item in data["groups"]}
    values = {item["key"]: item for item in data["values"]}
    for template in data["templates"]:
        lines = template["attribute_lines"]
        attribute_names = [line["attribute_key"] for line in lines]
        rules, final_suffix = split_name_format(
            template["name_format"], attribute_names
        )
        visible_lines = [
            (line, rules[normalize_token(line["attribute_key"])])
            for line in lines
            if rules[normalize_token(line["attribute_key"])]["mode"] != "hidden"
        ]
        leading_text = clean(visible_lines[0][1]["prefix"]) if visible_lines else ""
        group = groups[template["group_key"]]
        group_name = template["group_name_component"] or group["name"]
        source_model_name = clean(
            template["model_name_component"] or template["name"]
        )
        model_name = "" if template["suppress_model_name"] else source_model_name
        model_name = clean(f"{model_name} {leading_text}")
        base_name = name_with_group(group_name, model_name)
        sku_prefix = f"{group['code']}{template['model_code']}"

        line_values = {
            line["attribute_key"]: {
                item["value_key"]: item for item in line["values"]
            }
            for line in lines
        }
        for variant in template["variants"]:
            sku_parts = [sku_prefix]
            rendered_name = base_name
            deferred_markers = []
            displayed = 0
            previous_suffix = ""
            last_had_text = False
            for line, value_key in zip(lines, variant["value_keys"]):
                value = values[value_key]
                assignment = line_values[line["attribute_key"]][value_key]
                sku_component = clean(
                    assignment.get("sku_override") or value["sku_component"]
                )
                if sku_component != "—":
                    sku_parts.append(sku_component)

                rule = rules[normalize_token(line["attribute_key"])]
                if rule["mode"] == "hidden":
                    continue
                name_override = assignment.get("name_override")
                if clean(name_override) == "—":
                    name_component = ""
                else:
                    name_component = clean(name_override or value["name"])
                if "פתיחה" in normalize_token(line["attribute_key"]):
                    name_component, marker = split_direction_marker(name_component)
                    if marker:
                        deferred_markers.append(marker)
                if name_component:
                    if rendered_name and not displayed:
                        rendered_name += " "
                    elif displayed:
                        rendered_name += previous_suffix
                    rendered_name += name_component
                    displayed += 1
                    last_had_text = True
                else:
                    last_had_text = False
                previous_suffix = rule["suffix"]
            terminal_text = final_suffix
            if not terminal_text and last_had_text:
                terminal_text = previous_suffix
            rendered_name += terminal_text
            if deferred_markers:
                rendered_name += " " + " ".join(deferred_markers)
            rendered_name = clean(rendered_name)
            expected_name, marker = split_direction_marker(variant["name"])
            expected_name = clean(f"{expected_name} {marker}") if marker else expected_name
            rendered_sku = "".join(sku_parts)
            if rendered_sku != variant["sku"]:
                fail(
                    f"Rendered SKU mismatch for {variant['sku']}: {rendered_sku}"
                )
            if rendered_name != expected_name:
                fail(
                    f"Rendered name mismatch for {variant['sku']}: "
                    f"{rendered_name!r} != {expected_name!r}"
                )


def validate_filters(data):
    templates = {
        item["key"]: item
        for item in data["templates"]
        if item["group_key"] == FILTER_GROUP
    }
    if set(templates) != set(FILTER_TEMPLATE_COUNTS):
        fail(
            f"Expected four logical filter templates, found {sorted(templates)}"
        )

    all_variants = []
    for key, expected_count in FILTER_TEMPLATE_COUNTS.items():
        template = templates[key]
        if [line["attribute_key"] for line in template["attribute_lines"]] != [
            "קיבולת נפשות"
        ]:
            fail(f"{key} should use capacity as its only variant axis")
        if template["cartesian_count"] != expected_count:
            fail(f"Wrong Cartesian count in {key}")
        if len(template["variants"]) != expected_count:
            fail(f"Wrong active variant count in {key}")
        all_variants.extend(template["variants"])

    if len(all_variants) != 18:
        fail("The four filter templates should contain exactly 18 variants")
    forbidden_name_fragments = ("למרחב מוגן", 'לממ"מ', "בהתקנה")
    for variant in all_variants:
        if any(part in variant["name"] for part in forbidden_name_fragments):
            fail(f"A technical description remains in {variant['sku']}")
        if variant["sku"] not in {"400601", "400609"}:
            if "max_protected_area_m2" not in variant:
                fail(f"Protected area is missing from {variant['sku']}")
            if variant.get("installation_type") != "overhead":
                fail(f"Installation type is missing from {variant['sku']}")


def validate_white_mamad_wings(data):
    templates = {
        item["key"]: item
        for item in data["templates"]
        if item["key"] in WHITE_WING_TEMPLATE_COUNTS
    }
    if set(templates) != set(WHITE_WING_TEMPLATE_COUNTS):
        fail("The white MAMAD wings should use exactly two logical templates")
    for key, expected_count in WHITE_WING_TEMPLATE_COUNTS.items():
        template = templates[key]
        if len(template["variants"]) != expected_count:
            fail(f"Wrong wing variant count in {key}")
        if [line["attribute_key"] for line in template["attribute_lines"]] != [
            "צורת פתיחה לדלת", "פתח אור לדלת", "גובה כנף"
        ]:
            fail(f"Wrong wing axes in {key}")


def validate_source_workbook(data, workbook_path):
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError("openpyxl is required for workbook comparison") from error

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    source_skus = set()
    for sheet_name, sku_column in (("Groups", 40), ("Standalone", 1)):
        sheet = workbook[sheet_name]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            raw_sku = row[sku_column - 1] if len(row) >= sku_column else None
            if raw_sku is None:
                continue
            sku = str(raw_sku).strip()
            if sku.endswith(".0"):
                sku = sku[:-2]
            source_skus.add(sku)

    catalog_skus = {
        variant["sku"]
        for template in data["templates"]
        for variant in template["variants"]
    } | {item["sku"] for item in data["unique_items"]}
    if source_skus != catalog_skus:
        fail(
            "Workbook mismatch: "
            f"missing={sorted(source_skus - catalog_skus)[:10]}, "
            f"extra={sorted(catalog_skus - source_skus)[:10]}"
        )


def validate_idempotency(json_path):
    normalizer = Path(__file__).with_name("normalize_catalog.py")
    with tempfile.TemporaryDirectory() as directory:
        temp_path = Path(directory) / "catalog.json"
        payload_path = temp_path.with_name("catalog.json.gz.b64")
        temp_path.write_bytes(json_path.read_bytes())
        payload_path.write_bytes(
            json_path.with_name("catalog.json.gz.b64").read_bytes()
        )
        before = (temp_path.read_bytes(), payload_path.read_bytes())
        subprocess.run(
            [sys.executable, str(normalizer), str(temp_path)],
            check=True,
        )
        after = (temp_path.read_bytes(), payload_path.read_bytes())
        if before != after:
            fail("normalize_catalog.py is not idempotent")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("json_path", type=Path)
    parser.add_argument("--workbook", type=Path)
    args = parser.parse_args()

    data = load_catalog(args.json_path)
    skus = validate_references(data)
    validate_rendered_output(data)
    validate_filters(data)
    validate_white_mamad_wings(data)
    validate_idempotency(args.json_path)
    if args.workbook:
        validate_source_workbook(data, args.workbook)
    print(
        json.dumps(
            {
                "status": "ok",
                "products": len(skus),
                "templates": len(data["templates"]),
                "attributes": len(data["attributes"]),
                "values": len(data["values"]),
                "filter_templates": 4,
                "filter_variants": 18,
                "white_mamad_wing_templates": 2,
                "white_mamad_wing_variants": 38,
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
